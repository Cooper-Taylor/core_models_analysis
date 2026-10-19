#!/usr/bin/env python3
"""Reconcile the on-disk ModelSEED v2 model sets against the paper's supplement.

The manuscript (Faria et al., bioRxiv 2023.10.04.556561) states 5,540 core
models and 5,420 genome-scale models. ``/scratch/ctaylor/core_models_kegg2/``
holds 5,683. This script shows where the difference comes from, so the counts
quoted in reports/BASELINE_BIOMASS_TAXONOMY.md can be rechecked rather than
taken on trust.

It reads the two supplementary workbooks that carry genome manifests:

  Supplementary Table 4, sheet ``taxonomy``       the core manifest
  Supplementary Table 5, sheet ``Figure 6 data``  the genome-scale manifest

Both are public at <https://bioseed.mcs.anl.gov/~fliu/modelseed2/>; download
them into a directory of their own first, e.g.::

    mkdir -p /tmp/ms2_supp && cd /tmp/ms2_supp
    curl -fsSLO 'https://bioseed.mcs.anl.gov/~fliu/modelseed2/Supplementary%20Tables/Supplementary%20Table%204.xlsx'
    curl -fsSLO 'https://bioseed.mcs.anl.gov/~fliu/modelseed2/Supplementary%20Tables/Supplementary%20Table%205.xlsx'

Then, passing the paths as arguments and running from the repo root -- never
with the download directory as the working directory, and never by pointing an
interpreter at a file inside it. ``-P`` keeps the script's own directory off
``sys.path[0]`` so a stray ``json.py`` next to the data cannot be imported;
``-I`` would do that too but also drops user site-packages, where ``openpyxl``
usually lives::

    python3 -P scripts/verify_ms2_supplement_counts.py \
        '/tmp/ms2_supp/Supplementary Table 4.xlsx' \
        '/tmp/ms2_supp/Supplementary Table 5.xlsx'

Expected result, as of 2026-10-19:

    Supp. T4 taxonomy      5,539 genomes   (paper says 5,540)
    Supp. T5 Figure 6 data 5,419 genomes   (paper says 5,420)
    on disk, not in T4       145           144 Archaea + 1 Bacteria
    T4 - T5 (the drop)       120           118 Cyanobacteriota + 2 others

145 + 118 = 263, which is exactly the core-only genome count that
build_baseline_biomass_taxonomy.py reports. The 118 are the cyanobacteria the
paper deliberately excluded from the genome-scale study; the 145 are genomes
the shipped core_models.tar.gz carries beyond the manuscript's analysis set.

Requires ``openpyxl``.
"""
from __future__ import annotations

import argparse
import csv
import glob
import os
import re
import sys
from collections import Counter
from pathlib import Path

ANALYSIS_DIR = Path(os.environ.get("CORE_MODELS_ANALYSIS_DIR",
                                   "/scratch/ctaylor/core_models_analysis"))
CORE_MODELS_DIR = Path(os.environ.get("CORE_MODELS_DIR",
                                      "/scratch/ctaylor/core_models_kegg2"))


def sheet_ids(path: str, sheet: str, column: int = 0,
              strip_suffix: bool = False) -> set[str]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"{path}: no sheet {sheet!r} (has {wb.sheetnames})")
    rows = list(wb[sheet].iter_rows(values_only=True))
    out = set()
    for row in rows[1:]:
        if not row or row[column] is None:
            continue
        val = str(row[column]).strip()
        if strip_suffix:
            val = re.sub(r"\.RAST.*$", "", val)
        if val:
            out.add(val)
    print(f"  {Path(path).name} [{sheet}]: {len(rows) - 1} rows, "
          f"{len(out)} distinct ids")
    return out


def load_taxonomy() -> dict[str, dict]:
    tax: dict[str, dict] = {}
    for name in ("ncbi_taxonomy.csv", "ncbi_taxonomy_fill.csv"):
        path = ANALYSIS_DIR / "results" / name
        if not path.exists():
            continue
        with open(path) as fh:
            for row in csv.DictReader(fh):
                acc = row["assembly_accession"]
                if acc not in tax or tax[acc].get("status") != "resolved":
                    tax[acc] = row
    return tax


def breakdown(ids: set[str], tax: dict, label: str) -> None:
    print(f"\n{label} (n={len(ids)}):")
    for phylum, n in Counter((tax.get(i, {}).get("phylum") or "Unassigned")
                             for i in ids).most_common():
        print(f"   {phylum:26s} {n}")
    doms = Counter((tax.get(i, {}).get("superkingdom") or "Unassigned")
                   for i in ids)
    print(f"   by domain: {doms.most_common()}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("supp_table_4", help="Supplementary Table 4 .xlsx")
    ap.add_argument("supp_table_5", help="Supplementary Table 5 .xlsx")
    ap.add_argument("--core-models-dir", type=Path, default=CORE_MODELS_DIR)
    ap.add_argument("--ms2-gsm-dir", type=Path,
                    default=Path(os.environ.get(
                        "MS2_GSM_DIR", "/scratch/ctaylor/modelseed2_gs_models")),
                    help="the downloaded genome-scale model set, for the "
                         "on-disk (not supplement) GSM genome list")
    args = ap.parse_args()

    print("reading supplement manifests:")
    t4 = sheet_ids(args.supp_table_4, "taxonomy")
    t5 = sheet_ids(args.supp_table_5, "Figure 6 data", strip_suffix=True)
    disk = {Path(p).stem for p in
            glob.glob(str(args.core_models_dir / "GCF_*.json"))}

    print()
    print(f"on-disk core models            {len(disk):6d}")
    print(f"Supp. T4 taxonomy (core)       {len(t4):6d}   paper states 5,540")
    print(f"Supp. T5 Figure 6 (genome-sc.) {len(t5):6d}   paper states 5,420")
    print(f"  on disk, not in T4           {len(disk - t4):6d}")
    print(f"  in T4, not on disk           {len(t4 - disk):6d}")
    print(f"  T4 - T5 (dropped for GSM)    {len(t5 and (t4 - t5)):6d}")

    tax = load_taxonomy()
    breakdown(disk - t4, tax, "Shipped beyond the manuscript's core set")
    breakdown(t4 - t5, tax, "Dropped from the genome-scale study (per Supp. T5)")

    # The supplement sheet and the deposited models disagree slightly, so
    # close the loop against what is actually on disk rather than asserting
    # that 145 + 118 = 263 and hoping the arithmetic lands.
    gsm_disk: set[str] = set()
    for sub in ("gmm", "auxotrophy"):
        for path in glob.glob(str(args.ms2_gsm_dir / sub / "cobra" / "*.json.gz")):
            gsm_disk.add(re.sub(r"\.RAST.*$", "", Path(path).name.split(".json")[0]))
    if not gsm_disk:
        print(f"\n(no genome-scale models found under {args.ms2_gsm_dir}; "
              f"skipping the on-disk cross-check)")
        return 0

    core_only = disk - gsm_disk
    print(f"\non-disk genome-scale genomes   {len(gsm_disk):6d}   "
          f"(Supp. T5 lists {len(t5)})")
    print(f"on disk: core-only genomes     {len(core_only):6d}")
    print(f"  of which beyond Supp. T4     {len(core_only & (disk - t4)):6d}")
    print(f"  of which in T4 but not T5    {len(core_only & (t4 - t5)):6d}")
    stragglers = (t4 - t5) - core_only
    if stragglers:
        print(f"  listed as dropped by Supp. T5 yet PRESENT in the deposited "
              f"genome-scale set: {len(stragglers)}")
        for sid in sorted(stragglers):
            lin = tax.get(sid, {})
            print(f"     {sid}  {lin.get('phylum', '?')}  "
                  f"{lin.get('organism_name', '')}")
    breakdown(core_only, tax, "Core-only genomes, as deposited")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
